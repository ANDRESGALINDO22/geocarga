"""
Matriz de Riesgos - Proyecto GeoCarga (America's en carga)
Riesgo inherente (sin controles) vs. riesgo residual (con la respuesta propuesta).

Escala 1-5 para Probabilidad e Impacto.  Riesgo = Probabilidad x Impacto (1 a 25).
Equivalencia con el documento: Alta = 4, Media = 3.

Pestañas de la aplicación:
  1. Registro de riesgos      -> tabla general (solo consulta; doble clic para editar)
  2. Agregar riesgo           -> alta de un riesgo nuevo (nombre, categoría, P e I inherentes)
  3. Editar riesgo            -> modificar o eliminar un riesgo existente
  4. Controles y respuesta    -> control/respuesta, requisito relacionado y riesgo residual
  5. Mapas de calor           -> matrices 5x5 inherente y residual
  6. Comparativo              -> barras inherente vs residual

Exportación: archivo Excel (.xlsx) con tablas reales (filtrables y ordenables).
"""
import tkinter as tk
from tkinter import ttk, messagebox, filedialog

from matplotlib.figure import Figure
from matplotlib.patches import Rectangle
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg

# ----------------------------------------------------------------------------
# Niveles de riesgo
# ----------------------------------------------------------------------------
NIVELES = [
    (1, 4, "Bajo", "#66bb6a"),
    (5, 9, "Medio", "#ffee58"),
    (10, 14, "Alto", "#ffa726"),
    (15, 25, "Crítico", "#ef5350"),
]
COLOR_NIVEL = {n: c for _, _, n, c in NIVELES}
CATEGORIAS = ["Gestión", "Técnico", "Seguridad", "Operativo", "Datos", "Calendario"]


def nivel(score):
    for lo, hi, nombre, _ in NIVELES:
        if lo <= score <= hi:
            return nombre
    return "Bajo"


def calcular(rg):
    """Devuelve (inherente, residual, reducción %)."""
    inh = rg["pi"] * rg["ii"]
    res = rg["pr"] * rg["ir"]
    red = (1 - res / inh) * 100 if inh else 0
    return inh, res, red


# ----------------------------------------------------------------------------
# Datos iniciales: tabla "Riesgos del proyecto" del documento.
# pi/ii = probabilidad/impacto inherente; pr/ir = probabilidad/impacto residual
# (residual = estimación inicial; debe validarse con el equipo y el docente).
# ----------------------------------------------------------------------------
def r(id_, nombre, cat, pi, ii, control, req, pr, ir):
    return dict(id=id_, riesgo=nombre, categoria=cat, pi=pi, ii=ii,
                control=control, requisito=req, pr=pr, ir=ir)


RIESGOS_INICIALES = [
    r("R-01", "Alcance excesivo", "Gestión", 4, 4,
      "Priorizar MVP y aplicar control de cambios", "MoSCoW / H1", 2, 3),
    r("R-02", "Falta de usuario real", "Gestión", 4, 3,
      "Usar un representante, escenarios y revisiones frecuentes", "NF-03", 2, 3),
    r("R-03", "Poco acceso a datos o APIs", "Datos", 4, 4,
      "Utilizar datos ficticios y adaptadores simulados", "F-16", 3, 2),
    r("R-04", "Complejidad de IoT o nube", "Técnico", 3, 4,
      "Usar el teléfono como GPS; rastreadores dedicados en otra fase", "F-07", 2, 2),
    r("R-05", "Reglas de carga incorrectas", "Operativo", 3, 4,
      "Parámetros configurables y validación experta", "F-05, F-18", 2, 3),
    r("R-06", "Pérdida o exposición de datos", "Seguridad", 3, 4,
      "Datos anonimizados, permisos, respaldos y secretos protegidos", "NF-01, NF-02, NF-10", 2, 3),
    r("R-07", "Desbalance de trabajo", "Gestión", 3, 3,
      "Asignar responsables, revisar capacidad y rotar tareas", "NF-08", 2, 2),
    r("R-08", "Integración tardía", "Técnico", 3, 4,
      "Integrar en cada sprint y automatizar pruebas básicas", "NF-12", 2, 3),
    r("R-09", "Problemas de usabilidad", "Operativo", 3, 3,
      "Probar prototipos temprano con tareas representativas", "NF-03, NF-11", 2, 2),
    r("R-10", "Demora del calendario", "Calendario", 4, 3,
      "Reservar la última semana para estabilización y documentación", "H5", 3, 2),
    r("R-11", "GPS impreciso o desactivado", "Técnico", 3, 4,
      "Mostrar precisión, validar permiso y conservar última ubicación confiable", "F-07, NF-14", 3, 2),
    r("R-12", "Consumo de batería y datos", "Técnico", 3, 3,
      "Intervalo configurable y medición del consumo en pruebas", "NF-15", 2, 2),
    r("R-13", "Pérdida de conectividad", "Operativo", 4, 3,
      "Guardar posiciones localmente (IndexedDB) y sincronizar", "F-22, F-24", 2, 2),
    r("R-14", "Acceso indebido a ubicaciones", "Seguridad", 3, 4,
      "Canales privados, autorización por rol y retención limitada", "F-02, NF-16", 2, 3),
]


# ----------------------------------------------------------------------------
# Exportación a Excel con tablas reales
# ----------------------------------------------------------------------------
def exportar_excel(riesgos, ruta):
    """Genera un .xlsx con tablas de Excel (con filtros) en varias hojas."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    def fill(hex_):
        h = hex_.lstrip("#")
        return PatternFill("solid", start_color=h, end_color=h)

    borde = Side(style="thin", color="BBBBBB")
    caja = Border(left=borde, right=borde, top=borde, bottom=borde)
    centro = Alignment(horizontal="center", vertical="center", wrap_text=True)
    izq = Alignment(horizontal="left", vertical="center", wrap_text=True)

    def agregar_tabla(ws, nombre, ref):
        t = Table(displayName=nombre, ref=ref)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(t)

    def anchos(ws, valores):
        for i, w in enumerate(valores, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

    wb = Workbook()

    # ---- Hoja 1: Registro de riesgos ---------------------------------------
    ws = wb.active
    ws.title = "Registro de riesgos"
    enc = ["ID", "Riesgo", "Categoría",
           "Prob. inherente", "Impacto inherente", "Riesgo inherente", "Nivel inherente",
           "Control / Respuesta", "Requisito relacionado",
           "Prob. residual", "Impacto residual", "Riesgo residual", "Nivel residual",
           "Reducción (%)"]
    ws.append(enc)
    for rg in riesgos:
        inh, res, red = calcular(rg)
        ws.append([rg["id"], rg["riesgo"], rg["categoria"],
                   rg["pi"], rg["ii"], inh, nivel(inh),
                   rg["control"], rg["requisito"],
                   rg["pr"], rg["ir"], res, nivel(res),
                   round(red / 100, 4)])
    n = len(riesgos)
    ult = n + 1
    for fila in ws.iter_rows(min_row=2, max_row=ult, max_col=len(enc)):
        for c in fila:
            c.border = caja
            c.alignment = izq if c.column in (2, 8, 9) else centro
    for fila in range(2, ult + 1):
        for col in (7, 13):  # columnas de nivel
            c = ws.cell(row=fila, column=col)
            c.fill = fill(COLOR_NIVEL[c.value])
            c.font = Font(bold=True)
        ws.cell(row=fila, column=14).number_format = "0%"
    for c in ws[1]:
        c.alignment = centro
    ws.row_dimensions[1].height = 32
    anchos(ws, [8, 32, 13, 11, 11, 11, 12, 55, 20, 11, 11, 11, 12, 11])
    ws.freeze_panes = "C2"
    if n:
        agregar_tabla(ws, "RegistroRiesgos", f"A1:{get_column_letter(len(enc))}{ult}")

    # ---- Hoja 2: Resumen por nivel -----------------------------------------
    ws2 = wb.create_sheet("Resumen por nivel")
    ws2.append(["Nivel", "Rango (PxI)", "Riesgos inherentes", "Riesgos residuales"])
    for lo, hi, nombre, _ in NIVELES:
        ci = sum(1 for rg in riesgos if nivel(calcular(rg)[0]) == nombre)
        cr = sum(1 for rg in riesgos if nivel(calcular(rg)[1]) == nombre)
        ws2.append([nombre, f"{lo} a {hi}", ci, cr])
    for fila in ws2.iter_rows(min_row=2, max_row=5, max_col=4):
        for c in fila:
            c.border = caja
            c.alignment = centro
    for fila in range(2, 6):
        c = ws2.cell(row=fila, column=1)
        c.fill = fill(COLOR_NIVEL[c.value])
        c.font = Font(bold=True)
    anchos(ws2, [14, 14, 20, 20])
    agregar_tabla(ws2, "ResumenNiveles", "A1:D5")

    # ---- Hoja 3: Resumen por categoría -------------------------------------
    ws3 = wb.create_sheet("Resumen por categoría")
    ws3.append(["Categoría", "Cantidad", "Inherente promedio", "Residual promedio", "Reducción promedio"])
    cats = sorted({rg["categoria"] for rg in riesgos})
    for cat in cats:
        grupo = [rg for rg in riesgos if rg["categoria"] == cat]
        k = len(grupo)
        pi_ = sum(calcular(g)[0] for g in grupo) / k
        pr_ = sum(calcular(g)[1] for g in grupo) / k
        rd_ = sum(calcular(g)[2] for g in grupo) / k / 100
        ws3.append([cat, k, round(pi_, 1), round(pr_, 1), round(rd_, 4)])
    ult3 = len(cats) + 1
    for fila in ws3.iter_rows(min_row=2, max_row=ult3, max_col=5):
        for c in fila:
            c.border = caja
            c.alignment = centro
    for fila in range(2, ult3 + 1):
        ws3.cell(row=fila, column=5).number_format = "0%"
    anchos(ws3, [16, 11, 20, 20, 20])
    if cats:
        agregar_tabla(ws3, "ResumenCategorias", f"A1:E{ult3}")

    # ---- Hoja 4: Matrices 5x5 (inherente y residual) -----------------------
    ws4 = wb.create_sheet("Mapas de calor")

    def dibujar_matriz(fila0, titulo, kp, ki):
        ws4.cell(row=fila0, column=1, value=titulo).font = Font(bold=True, size=12)
        c0 = ws4.cell(row=fila0 + 1, column=1, value="Prob. \\ Impacto")
        c0.font = Font(bold=True)
        c0.alignment = centro
        c0.border = caja
        for i in range(1, 6):
            c = ws4.cell(row=fila0 + 1, column=i + 1, value=f"Impacto {i}")
            c.font = Font(bold=True)
            c.alignment = centro
            c.border = caja
        for p in range(5, 0, -1):  # probabilidad 5 arriba
            f = fila0 + 2 + (5 - p)
            c = ws4.cell(row=f, column=1, value=f"Prob. {p}")
            c.font = Font(bold=True)
            c.alignment = centro
            c.border = caja
            ws4.row_dimensions[f].height = 42
            for i in range(1, 6):
                ids = [x["id"] for x in riesgos if x[kp] == p and x[ki] == i]
                celda = ws4.cell(row=f, column=i + 1,
                                 value=(", ".join(ids) if ids else f"({p * i})"))
                celda.fill = fill(COLOR_NIVEL[nivel(p * i)])
                celda.alignment = centro
                celda.border = caja
                celda.font = Font(bold=bool(ids), color="000000" if ids else "666666")

    dibujar_matriz(1, "Riesgo INHERENTE (sin controles)", "pi", "ii")
    dibujar_matriz(10, "Riesgo RESIDUAL (con controles)", "pr", "ir")
    ws4.cell(row=18, column=1, value="Cada celda muestra los ID de los riesgos; "
                                     "entre paréntesis el valor PxI cuando está vacía.").font = Font(italic=True)
    anchos(ws4, [16, 18, 18, 18, 18, 18])

    wb.save(ruta)


# ----------------------------------------------------------------------------
# Aplicación
# ----------------------------------------------------------------------------
class MatrizRiesgoApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Matriz de Riesgos GeoCarga - Inherente vs Residual")
        self.root.geometry("1250x780")
        self.riesgos = [dict(x) for x in RIESGOS_INICIALES]

        self.crear_widgets()
        self.refrescar_todo()

    # ------------------------------------------------------------------ UI
    def crear_widgets(self):
        main = ttk.Frame(self.root, padding=8)
        main.pack(fill=tk.BOTH, expand=True)

        top = ttk.Frame(main)
        top.pack(fill=tk.X, pady=(0, 6))
        ttk.Button(top, text="Restaurar datos del documento", command=self.restaurar).pack(side=tk.LEFT)
        ttk.Button(top, text="Exportar a Excel", command=self.exportar).pack(side=tk.LEFT, padx=6)
        self.lbl_resumen = ttk.Label(top, text="", font=("Segoe UI", 9, "bold"))
        self.lbl_resumen.pack(side=tk.RIGHT)

        self.nb = ttk.Notebook(main)
        self.nb.pack(fill=tk.BOTH, expand=True)

        self.tab_reg = ttk.Frame(self.nb)
        self.tab_add = ttk.Frame(self.nb, padding=20)
        self.tab_edit = ttk.Frame(self.nb, padding=20)
        self.tab_ctrl = ttk.Frame(self.nb, padding=20)
        self.tab_mapas = ttk.Frame(self.nb)
        self.tab_comp = ttk.Frame(self.nb)
        self.nb.add(self.tab_reg, text="Registro de riesgos")
        self.nb.add(self.tab_add, text="Agregar riesgo")
        self.nb.add(self.tab_edit, text="Editar riesgo")
        self.nb.add(self.tab_ctrl, text="Controles y respuesta")
        self.nb.add(self.tab_mapas, text="Mapas de calor")
        self.nb.add(self.tab_comp, text="Comparativo")

        self._tab_registro()
        self._tab_agregar()
        self._tab_editar()
        self._tab_controles()

        self.fig_mapas = Figure(figsize=(11, 5), dpi=100)
        self.canvas_mapas = FigureCanvasTkAgg(self.fig_mapas, master=self.tab_mapas)
        self.canvas_mapas.get_tk_widget().pack(fill=tk.BOTH, expand=True)

        self.fig_comp = Figure(figsize=(11, 5), dpi=100)
        self.canvas_comp = FigureCanvasTkAgg(self.fig_comp, master=self.tab_comp)
        self.canvas_comp.get_tk_widget().pack(fill=tk.BOTH, expand=True)

    # ---- helpers de formulario ---------------------------------------------
    @staticmethod
    def _spin(parent, valor):
        var = tk.StringVar(value=str(valor))
        sp = ttk.Spinbox(parent, from_=1, to=5, width=5, textvariable=var, justify="center")
        return sp, var

    @staticmethod
    def _titulo(parent, texto, sub=None):
        ttk.Label(parent, text=texto, font=("Segoe UI", 13, "bold")).grid(
            row=0, column=0, columnspan=4, sticky=tk.W)
        if sub:
            ttk.Label(parent, text=sub, foreground="#555", wraplength=760, justify=tk.LEFT).grid(
                row=1, column=0, columnspan=4, sticky=tk.W, pady=(2, 14))

    @staticmethod
    def _vista_previa(parent, fila, var_p, var_i, titulo):
        """Etiqueta que muestra P x I y nivel en vivo."""
        lbl = tk.Label(parent, text="", font=("Segoe UI", 10, "bold"), width=38, pady=6)
        lbl.grid(row=fila, column=0, columnspan=4, sticky=tk.W, pady=8)

        def actualizar(*_):
            try:
                p, i = int(var_p.get()), int(var_i.get())
                if not (1 <= p <= 5 and 1 <= i <= 5):
                    raise ValueError
                lbl.config(text=f"{titulo}: {p} × {i} = {p * i}  ({nivel(p * i)})",
                           bg=COLOR_NIVEL[nivel(p * i)])
            except ValueError:
                lbl.config(text=f"{titulo}: valores no válidos", bg="#eeeeee")

        var_p.trace_add("write", actualizar)
        var_i.trace_add("write", actualizar)
        actualizar()

    def _opciones_riesgos(self):
        return [f"{x['id']} - {x['riesgo']}" for x in self.riesgos]

    def _riesgo_de_combo(self, combo):
        txt = combo.get()
        if not txt:
            return None
        id_ = txt.split(" - ")[0]
        for rg in self.riesgos:
            if rg["id"] == id_:
                return rg
        return None

    # ---- Pestaña: Registro ---------------------------------------------------
    def _tab_registro(self):
        cols = ("id", "riesgo", "cat", "pi", "ii", "inh", "niv_inh",
                "control", "pr", "ir", "res", "niv_res", "red")
        heads = ("ID", "Riesgo", "Categoría", "P", "I", "Inherente", "Nivel inh.",
                 "Control / Respuesta", "P res.", "I res.", "Residual", "Nivel res.", "Reducción")
        widths = (50, 190, 80, 30, 30, 70, 70, 330, 50, 50, 65, 75, 75)

        frame = ttk.Frame(self.tab_reg)
        frame.pack(fill=tk.BOTH, expand=True)
        self.tree = ttk.Treeview(frame, columns=cols, show="headings", selectmode="browse")
        for c, h, w in zip(cols, heads, widths):
            self.tree.heading(c, text=h)
            self.tree.column(c, width=w, anchor=tk.W if c in ("riesgo", "control") else tk.CENTER)
        for nombre, color in COLOR_NIVEL.items():
            self.tree.tag_configure(nombre, background=color)
        sb = ttk.Scrollbar(frame, orient=tk.VERTICAL, command=self.tree.yview)
        self.tree.configure(yscroll=sb.set)
        sb.pack(side=tk.RIGHT, fill=tk.Y)
        self.tree.pack(fill=tk.BOTH, expand=True)
        self.tree.bind("<Double-1>", self.doble_clic_registro)

        ttk.Label(self.tab_reg,
                  text="Color de fila = nivel del riesgo RESIDUAL.  "
                       "Doble clic en una fila para editarla.",
                  foreground="#555").pack(anchor=tk.W, padx=4, pady=2)

    def doble_clic_registro(self, _evt):
        sel = self.tree.selection()
        if not sel:
            return
        self.edit_combo.set(self._opciones_riesgos()[self.tree.index(sel[0])])
        self.cargar_edicion()
        self.nb.select(self.tab_edit)

    # ---- Pestaña: Agregar ----------------------------------------------------
    def _tab_agregar(self):
        f = self.tab_add
        self._titulo(f, "Agregar un riesgo nuevo",
                     "Registra el riesgo con su probabilidad e impacto INHERENTES (sin controles). "
                     "Luego define su control y su riesgo residual en la pestaña "
                     "«Controles y respuesta».")

        ttk.Label(f, text="Nombre del riesgo:").grid(row=2, column=0, sticky=tk.E, pady=5)
        self.a_nombre = ttk.Entry(f, width=45)
        self.a_nombre.grid(row=2, column=1, columnspan=3, sticky=tk.W, padx=8)

        ttk.Label(f, text="Categoría:").grid(row=3, column=0, sticky=tk.E, pady=5)
        self.a_cat = ttk.Combobox(f, values=CATEGORIAS, width=22, state="readonly")
        self.a_cat.current(0)
        self.a_cat.grid(row=3, column=1, sticky=tk.W, padx=8)

        ttk.Label(f, text="Probabilidad inherente (1-5):").grid(row=4, column=0, sticky=tk.E, pady=5)
        self.a_pi, self.a_pi_v = self._spin(f, 3)
        self.a_pi.grid(row=4, column=1, sticky=tk.W, padx=8)

        ttk.Label(f, text="Impacto inherente (1-5):").grid(row=5, column=0, sticky=tk.E, pady=5)
        self.a_ii, self.a_ii_v = self._spin(f, 3)
        self.a_ii.grid(row=5, column=1, sticky=tk.W, padx=8)

        self._vista_previa(f, 6, self.a_pi_v, self.a_ii_v, "Riesgo inherente")

        bt = ttk.Frame(f)
        bt.grid(row=7, column=0, columnspan=4, sticky=tk.W, pady=10)
        ttk.Button(bt, text="Agregar riesgo", command=self.agregar).pack(side=tk.LEFT, padx=(0, 8))
        ttk.Button(bt, text="Limpiar", command=self.limpiar_agregar).pack(side=tk.LEFT)

    # ---- Pestaña: Editar -----------------------------------------------------
    def _tab_editar(self):
        f = self.tab_edit
        self._titulo(f, "Editar un riesgo existente",
                     "Selecciona un riesgo para modificar su nombre, categoría y su valoración inherente. "
                     "El control y el riesgo residual se editan en la pestaña «Controles y respuesta».")

        ttk.Label(f, text="Riesgo a editar:").grid(row=2, column=0, sticky=tk.E, pady=5)
        self.edit_combo = ttk.Combobox(f, width=48, state="readonly")
        self.edit_combo.grid(row=2, column=1, columnspan=3, sticky=tk.W, padx=8)
        self.edit_combo.bind("<<ComboboxSelected>>", lambda e: self.cargar_edicion())

        ttk.Label(f, text="Nombre del riesgo:").grid(row=3, column=0, sticky=tk.E, pady=5)
        self.e_nombre = ttk.Entry(f, width=45)
        self.e_nombre.grid(row=3, column=1, columnspan=3, sticky=tk.W, padx=8)

        ttk.Label(f, text="Categoría:").grid(row=4, column=0, sticky=tk.E, pady=5)
        self.e_cat = ttk.Combobox(f, values=CATEGORIAS, width=22, state="readonly")
        self.e_cat.grid(row=4, column=1, sticky=tk.W, padx=8)

        ttk.Label(f, text="Probabilidad inherente (1-5):").grid(row=5, column=0, sticky=tk.E, pady=5)
        self.e_pi, self.e_pi_v = self._spin(f, 3)
        self.e_pi.grid(row=5, column=1, sticky=tk.W, padx=8)

        ttk.Label(f, text="Impacto inherente (1-5):").grid(row=6, column=0, sticky=tk.E, pady=5)
        self.e_ii, self.e_ii_v = self._spin(f, 3)
        self.e_ii.grid(row=6, column=1, sticky=tk.W, padx=8)

        self._vista_previa(f, 7, self.e_pi_v, self.e_ii_v, "Riesgo inherente")

        bt = ttk.Frame(f)
        bt.grid(row=8, column=0, columnspan=4, sticky=tk.W, pady=10)
        ttk.Button(bt, text="Guardar cambios", command=self.guardar_edicion).pack(side=tk.LEFT, padx=(0, 8))
        ttk.Button(bt, text="Eliminar riesgo", command=self.eliminar).pack(side=tk.LEFT)

    # ---- Pestaña: Controles y respuesta --------------------------------------
    def _tab_controles(self):
        f = self.tab_ctrl
        self._titulo(f, "Controles, respuesta y riesgo residual",
                     "Selecciona un riesgo, describe el control o respuesta que se aplicará y "
                     "valora la probabilidad e impacto que quedan DESPUÉS del control (residual).")

        ttk.Label(f, text="Riesgo:").grid(row=2, column=0, sticky=tk.E, pady=5)
        self.c_combo = ttk.Combobox(f, width=48, state="readonly")
        self.c_combo.grid(row=2, column=1, columnspan=3, sticky=tk.W, padx=8)
        self.c_combo.bind("<<ComboboxSelected>>", lambda e: self.cargar_control())

        self.c_inh = tk.Label(f, text="Riesgo inherente actual: (selecciona un riesgo)",
                              font=("Segoe UI", 10, "bold"), width=48, pady=6, bg="#eeeeee")
        self.c_inh.grid(row=3, column=0, columnspan=4, sticky=tk.W, pady=6)

        ttk.Label(f, text="Control / respuesta:").grid(row=4, column=0, sticky=tk.NE, pady=5)
        self.c_texto = tk.Text(f, width=58, height=4, wrap=tk.WORD, font=("Segoe UI", 9))
        self.c_texto.grid(row=4, column=1, columnspan=3, sticky=tk.W, padx=8)

        ttk.Label(f, text="Requisito relacionado:").grid(row=5, column=0, sticky=tk.E, pady=5)
        self.c_req = ttk.Entry(f, width=30)
        self.c_req.grid(row=5, column=1, columnspan=2, sticky=tk.W, padx=8)

        ttk.Label(f, text="Probabilidad residual (1-5):").grid(row=6, column=0, sticky=tk.E, pady=5)
        self.c_pr, self.c_pr_v = self._spin(f, 2)
        self.c_pr.grid(row=6, column=1, sticky=tk.W, padx=8)

        ttk.Label(f, text="Impacto residual (1-5):").grid(row=7, column=0, sticky=tk.E, pady=5)
        self.c_ir, self.c_ir_v = self._spin(f, 2)
        self.c_ir.grid(row=7, column=1, sticky=tk.W, padx=8)

        self._vista_previa(f, 8, self.c_pr_v, self.c_ir_v, "Riesgo residual")

        bt = ttk.Frame(f)
        bt.grid(row=9, column=0, columnspan=4, sticky=tk.W, pady=10)
        ttk.Button(bt, text="Guardar control y residual", command=self.guardar_control).pack(side=tk.LEFT)

    # ------------------------------------------------------------- acciones
    @staticmethod
    def _leer_1a5(*vars_):
        try:
            vals = [int(v.get()) for v in vars_]
        except ValueError:
            return None
        return vals if all(1 <= x <= 5 for x in vals) else None

    def siguiente_id(self):
        nums = [int(x["id"].split("-")[1]) for x in self.riesgos
                if x["id"].startswith("R-") and x["id"].split("-")[1].isdigit()]
        return f"R-{(max(nums) + 1) if nums else 1:02d}"

    # Agregar
    def agregar(self):
        nombre = self.a_nombre.get().strip()
        if not nombre:
            messagebox.showerror("Error", "El nombre del riesgo es obligatorio.")
            return
        vals = self._leer_1a5(self.a_pi_v, self.a_ii_v)
        if vals is None:
            messagebox.showerror("Error", "Probabilidad e impacto deben ser enteros entre 1 y 5.")
            return
        pi, ii = vals
        nuevo = dict(id=self.siguiente_id(), riesgo=nombre, categoria=self.a_cat.get() or "Gestión",
                     pi=pi, ii=ii, control="Sin control definido", requisito="",
                     pr=pi, ir=ii)  # sin control: residual = inherente
        self.riesgos.append(nuevo)
        self.limpiar_agregar()
        self.refrescar_todo()
        self.c_combo.set(f"{nuevo['id']} - {nuevo['riesgo']}")
        self.cargar_control()
        if messagebox.askyesno("Riesgo agregado",
                               f"Se creó {nuevo['id']}. Su riesgo residual es igual al inherente "
                               "porque aún no tiene control.\n\n¿Definir ahora su control y respuesta?"):
            self.nb.select(self.tab_ctrl)

    def limpiar_agregar(self):
        self.a_nombre.delete(0, tk.END)
        self.a_cat.current(0)
        self.a_pi_v.set("3")
        self.a_ii_v.set("3")

    # Editar
    def cargar_edicion(self):
        rg = self._riesgo_de_combo(self.edit_combo)
        if not rg:
            return
        self.e_nombre.delete(0, tk.END)
        self.e_nombre.insert(0, rg["riesgo"])
        self.e_cat.set(rg["categoria"])
        self.e_pi_v.set(str(rg["pi"]))
        self.e_ii_v.set(str(rg["ii"]))

    def guardar_edicion(self):
        rg = self._riesgo_de_combo(self.edit_combo)
        if not rg:
            messagebox.showinfo("Info", "Selecciona primero el riesgo que deseas editar.")
            return
        nombre = self.e_nombre.get().strip()
        if not nombre:
            messagebox.showerror("Error", "El nombre del riesgo es obligatorio.")
            return
        vals = self._leer_1a5(self.e_pi_v, self.e_ii_v)
        if vals is None:
            messagebox.showerror("Error", "Probabilidad e impacto deben ser enteros entre 1 y 5.")
            return
        rg.update(riesgo=nombre, categoria=self.e_cat.get() or rg["categoria"], pi=vals[0], ii=vals[1])
        id_ = rg["id"]
        self.refrescar_todo()
        self.edit_combo.set(f"{id_} - {nombre}")
        if self._riesgo_de_combo(self.c_combo) is rg:
            self.c_combo.set(f"{id_} - {nombre}")
            self.cargar_control()
        if rg["pr"] * rg["ir"] > rg["pi"] * rg["ii"]:
            messagebox.showwarning(
                "Revisar residual",
                "Con este cambio el riesgo residual quedó MAYOR que el inherente.\n"
                "Ajusta el residual en la pestaña «Controles y respuesta».")
        else:
            messagebox.showinfo("Guardado", f"{id_} actualizado.")

    def eliminar(self):
        rg = self._riesgo_de_combo(self.edit_combo)
        if not rg:
            messagebox.showinfo("Info", "Selecciona primero el riesgo que deseas eliminar.")
            return
        if messagebox.askyesno("Eliminar", f"¿Eliminar {rg['id']} - {rg['riesgo']}?"):
            self.riesgos.remove(rg)
            self.edit_combo.set("")
            self.c_combo.set("")
            self.e_nombre.delete(0, tk.END)
            self.c_texto.delete("1.0", tk.END)
            self.c_req.delete(0, tk.END)
            self.refrescar_todo()

    # Controles
    def cargar_control(self):
        rg = self._riesgo_de_combo(self.c_combo)
        if not rg:
            return
        inh = rg["pi"] * rg["ii"]
        self.c_inh.config(text=f"Riesgo inherente actual: {rg['pi']} × {rg['ii']} = {inh}  ({nivel(inh)})",
                          bg=COLOR_NIVEL[nivel(inh)])
        self.c_texto.delete("1.0", tk.END)
        self.c_texto.insert("1.0", rg["control"])
        self.c_req.delete(0, tk.END)
        self.c_req.insert(0, rg["requisito"])
        self.c_pr_v.set(str(rg["pr"]))
        self.c_ir_v.set(str(rg["ir"]))

    def guardar_control(self):
        rg = self._riesgo_de_combo(self.c_combo)
        if not rg:
            messagebox.showinfo("Info", "Selecciona primero un riesgo.")
            return
        vals = self._leer_1a5(self.c_pr_v, self.c_ir_v)
        if vals is None:
            messagebox.showerror("Error", "Probabilidad e impacto deben ser enteros entre 1 y 5.")
            return
        pr, ir = vals
        if pr * ir > rg["pi"] * rg["ii"]:
            if not messagebox.askyesno(
                    "Revisar",
                    "El riesgo residual es MAYOR que el inherente.\n"
                    "Normalmente los controles reducen el riesgo. ¿Guardar de todos modos?"):
                return
        rg.update(control=self.c_texto.get("1.0", tk.END).strip(),
                  requisito=self.c_req.get().strip(), pr=pr, ir=ir)
        self.refrescar_todo()
        messagebox.showinfo("Guardado", f"Control y residual de {rg['id']} actualizados.")

    # Generales
    def restaurar(self):
        if messagebox.askyesno("Restaurar", "Se perderán los cambios. ¿Continuar?"):
            self.riesgos = [dict(x) for x in RIESGOS_INICIALES]
            for cb in (self.edit_combo, self.c_combo):
                cb.set("")
            self.refrescar_todo()

    def exportar(self):
        if not self.riesgos:
            messagebox.showinfo("Info", "No hay riesgos para exportar.")
            return
        ruta = filedialog.asksaveasfilename(
            defaultextension=".xlsx", filetypes=[("Libro de Excel", "*.xlsx")],
            initialfile="matriz_riesgos_geocarga.xlsx")
        if not ruta:
            return
        try:
            exportar_excel(self.riesgos, ruta)
        except PermissionError:
            messagebox.showerror("Error", "No se pudo guardar. Cierra el archivo si lo tienes abierto en Excel.")
            return
        messagebox.showinfo("Exportado", f"Archivo guardado en:\n{ruta}")

    # ------------------------------------------------------------ refresco
    def refrescar_todo(self):
        opciones = self._opciones_riesgos()
        self.edit_combo["values"] = opciones
        self.c_combo["values"] = opciones
        self.actualizar_tabla()
        self.dibujar_mapas()
        self.dibujar_comparativo()
        self.actualizar_resumen()

    def actualizar_tabla(self):
        self.tree.delete(*self.tree.get_children())
        for rg in self.riesgos:
            inh, res, red = calcular(rg)
            self.tree.insert("", tk.END, tags=(nivel(res),), values=(
                rg["id"], rg["riesgo"], rg["categoria"], rg["pi"], rg["ii"], inh, nivel(inh),
                rg["control"], rg["pr"], rg["ir"], res, nivel(res), f"{red:.0f}%"))

    def actualizar_resumen(self):
        if not self.riesgos:
            self.lbl_resumen.config(text="Sin riesgos")
            return
        niv_i = [nivel(calcular(x)[0]) for x in self.riesgos]
        niv_r = [nivel(calcular(x)[1]) for x in self.riesgos]
        red = sum(calcular(x)[2] for x in self.riesgos) / len(self.riesgos)
        self.lbl_resumen.config(
            text=f"Inherente: {niv_i.count('Crítico')} críticos / {niv_i.count('Alto')} altos   |   "
                 f"Residual: {niv_r.count('Crítico')} críticos / {niv_r.count('Alto')} altos   |   "
                 f"Reducción media: {red:.0f}%")

    # ------------------------------------------------------------- gráficos
    def _matriz(self, ax, kp, ki, titulo):
        ax.set_xlim(0, 5)
        ax.set_ylim(0, 5)
        for p in range(1, 6):
            for i in range(1, 6):
                color = COLOR_NIVEL[nivel(p * i)]
                ax.add_patch(Rectangle((i - 1, p - 1), 1, 1, facecolor=color,
                                       edgecolor="white", linewidth=2, alpha=0.85))
                ax.text(i - 0.93, p - 0.2, str(p * i), fontsize=7, color="#444")
                ids = [rg["id"].replace("R-", "") for rg in self.riesgos
                       if rg[kp] == p and rg[ki] == i]
                if ids:
                    lineas = [", ".join(ids[j:j + 3]) for j in range(0, len(ids), 3)]
                    ax.text(i - 0.5, p - 0.5, "\n".join(lineas), ha="center", va="center",
                            fontsize=9, fontweight="bold")
        ax.set_xticks([x + 0.5 for x in range(5)])
        ax.set_xticklabels(["1\nMuy bajo", "2\nBajo", "3\nMedio", "4\nAlto", "5\nMuy alto"], fontsize=8)
        ax.set_yticks([x + 0.5 for x in range(5)])
        ax.set_yticklabels(["1", "2", "3", "4", "5"], fontsize=8)
        ax.set_xlabel("Impacto")
        ax.set_ylabel("Probabilidad")
        ax.set_title(titulo, fontweight="bold")
        ax.set_aspect("equal")

    def dibujar_mapas(self):
        self.fig_mapas.clear()
        ax1 = self.fig_mapas.add_subplot(1, 2, 1)
        ax2 = self.fig_mapas.add_subplot(1, 2, 2)
        self._matriz(ax1, "pi", "ii", "Riesgo INHERENTE (sin controles)")
        self._matriz(ax2, "pr", "ir", "Riesgo RESIDUAL (con controles)")
        self.fig_mapas.text(0.5, 0.01, "Los números en cada celda corresponden al ID del riesgo (R-xx).",
                            ha="center", fontsize=8, color="#555")
        self.fig_mapas.tight_layout(rect=(0, 0.03, 1, 1))
        self.canvas_mapas.draw()

    def dibujar_comparativo(self):
        self.fig_comp.clear()
        ax = self.fig_comp.add_subplot(111)
        if self.riesgos:
            ids = [rg["id"] for rg in self.riesgos]
            inh = [calcular(rg)[0] for rg in self.riesgos]
            res = [calcular(rg)[1] for rg in self.riesgos]
            x = range(len(ids))
            w = 0.38
            ax.bar([i - w / 2 for i in x], inh, w, label="Inherente", color="#c62828")
            ax.bar([i + w / 2 for i in x], res, w, label="Residual", color="#2e7d32")
            for i, (a, b) in enumerate(zip(inh, res)):
                ax.text(i - w / 2, a + 0.3, str(a), ha="center", fontsize=8)
                ax.text(i + w / 2, b + 0.3, str(b), ha="center", fontsize=8)
            ax.axhline(15, color="#ef5350", linestyle="--", linewidth=1)
            ax.axhline(10, color="#ffa726", linestyle="--", linewidth=1)
            ax.text(len(ids) - 0.5, 15.3, "Crítico ≥ 15", color="#ef5350", ha="right", fontsize=8)
            ax.text(len(ids) - 0.5, 10.3, "Alto ≥ 10", color="#fb8c00", ha="right", fontsize=8)
            ax.set_xticks(list(x))
            ax.set_xticklabels(ids, rotation=45, fontsize=8)
            ax.set_ylim(0, 27)
            ax.legend()
        ax.set_ylabel("Nivel de riesgo (P × I)")
        ax.set_title("Riesgo inherente vs. residual por riesgo", fontweight="bold")
        self.fig_comp.tight_layout()
        self.canvas_comp.draw()


if __name__ == "__main__":
    root = tk.Tk()
    app = MatrizRiesgoApp(root)
    root.mainloop()